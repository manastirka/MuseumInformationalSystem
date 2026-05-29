import os

os.environ.setdefault('FLASK_ENV', 'testing')
os.environ.setdefault('SECRET_KEY', 'test-secret')
os.environ.setdefault('REDIS_URL', '')
os.environ.setdefault('SESSION_TYPE', 'filesystem')
os.environ.setdefault('SESSION_FILE_DIR', '/tmp/museum-test-b-timesheet-export')

from contextlib import contextmanager
from decimal import Decimal
from unittest import mock

import pytest

import timesheet_word_export
from docx import Document


# A normal full-time work day uses 8h. The reproduction targets a month
# (May 2025, 31 days) where some days carry half-hour worked values and
# some carry whole-day paid absence (vacation / holiday / sick).

HEADER = {
    'id': 1,
    'employee_name': 'Петар Петровић',
    'month': 5,
    'year': 2025,
    'organization_unit': 'Минералошка збирка',
    'position': 'Кустос',
    'special_tasks': None,
    'extraordinary_tasks': None,
    'duties_summary': None,
    'created_at': None,
}


def _day(day, **kw):
    base = {
        'day': day,
        'work_in_museum': Decimal('0'),
        'work_outside': Decimal('0'),
        'vacation': Decimal('0'),
        'public_holiday': Decimal('0'),
        'paid_leave': Decimal('0'),
        'other_leave': Decimal('0'),
        'sick_leave_lt30': Decimal('0'),
        'sick_leave_gte30': Decimal('0'),
    }
    for k, v in kw.items():
        base[k] = Decimal(str(v))
    return base


# Daily rows:
#  day 1 (Thu): 4.5h worked in museum  -> must render "4.5", not "4"
#  day 2 (Fri): 7.5h worked in museum  -> must render "7.5", not "7"
#  day 5 (Mon): 8h vacation (paid absence, must NOT count toward worked grand total)
#  day 6 (Tue): 8h public holiday
#  day 7 (Wed): 8h sick_leave_lt30
DAILY_ROWS = [
    _day(1, work_in_museum='4.5'),
    _day(2, work_in_museum='7.5'),
    _day(5, vacation='8'),
    _day(6, public_holiday='8'),
    _day(7, sick_leave_lt30='8'),
]

# Worked hours only: 4.5 + 7.5 = 12.0
WORKED_TOTAL = 12.0


class _FakeCursor:
    def __init__(self):
        self._result = None

    def execute(self, sql, params=None):
        if 'FROM timesheet_reports' in sql:
            self._result = ('one', HEADER)
        elif 'FROM timesheet_report_days' in sql:
            self._result = ('all', list(DAILY_ROWS))
        else:
            self._result = ('one', None)

    def fetchone(self):
        kind, data = self._result
        return data if kind == 'one' else None

    def fetchall(self):
        kind, data = self._result
        return data if kind == 'all' else []

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class _FakeConn:
    def cursor(self):
        return _FakeCursor()

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


@contextmanager
def _fake_connect(*a, **k):
    yield _FakeConn()


def _render(tmp_path):
    with mock.patch.object(timesheet_word_export.psycopg, 'connect', _fake_connect), \
            mock.patch.object(timesheet_word_export.os.path, 'dirname', return_value=str(tmp_path)):
        path = timesheet_word_export.generate_word_document(1, 'postgresql://x/y')
    return path


def _main_table_cells(path):
    doc = Document(path)
    table = doc.tables[0]
    return table


def test_half_hour_cell_values_not_truncated(tmp_path):
    path = _render(tmp_path)
    table = _main_table_cells(path)
    # Row 4 = 'Радни сати у музеју' (work_in_museum). Col == day number.
    assert table.cell(4, 1).text.strip() == '4.5'
    assert table.cell(4, 2).text.strip() == '7.5'


def test_half_hour_row_total_not_truncated(tmp_path):
    path = _render(tmp_path)
    table = _main_table_cells(path)
    # Row total for work_in_museum row is in col 32: 4.5 + 7.5 = 12.0
    assert table.cell(4, 32).text.strip() == '12'


def test_grand_total_excludes_leave_and_sick(tmp_path):
    path = _render(tmp_path)
    table = _main_table_cells(path)
    # Grand total cell is (12, 32). Must equal worked hours only (12),
    # NOT 12 + 8 vacation + 8 holiday + 8 sick = 36.
    assert table.cell(12, 32).text.strip() == '12'


def test_whole_hours_still_render_without_decimal(tmp_path):
    path = _render(tmp_path)
    table = _main_table_cells(path)
    # Vacation row is row 6; day 5 has 8h vacation -> '8' (no '8.0').
    assert table.cell(6, 5).text.strip() == '8'
